import sys
from pptx import Presentation

def extract_text(file_path):
    prs = Presentation(file_path)
    text_content = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_content.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_data = [cell.text for cell in row.cells]
                    text_content.append(" | ".join(row_data))
    return "\n".join(text_content)

if __name__ == "__main__":
    file_path = "301374966800003_20260103T202406_INV-000030.pptx"
    print(extract_text(file_path))
